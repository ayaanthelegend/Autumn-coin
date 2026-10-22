"""Update slides 17-19 in Autumn_Coin_IYPT_2026_Full_22Slides.pptx with C++ SPH results."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT_DIR = Path(__file__).resolve().parent.parent
PPTX_PATH = ROOT_DIR / "Autumn_Coin_IYPT_2026_Full_22Slides.pptx"
VAL_PLOT = ROOT_DIR / "results_sph" / "validation" / "validation_ladder_summary.png"
COMP_PLOT = ROOT_DIR / "results_sph" / "comparison" / "sph_vs_ode_comparison.png"
FRAMES_DIR = ROOT_DIR / "results_sph" / "drops" / "frames"

def update_slides():
    print("[Slide Updater] Loading presentation...")
    prs = Presentation(str(PPTX_PATH))

    # --- SLIDE 17: Validation Ladder ---
    slide17 = prs.slides[16]
    print("[Slide Updater] Updating Slide 17 (Validation Ladder)...")
    if VAL_PLOT.exists():
        # Place image on right half
        left = Inches(6.5)
        top = Inches(1.4)
        width = Inches(6.5)
        slide17.shapes.add_picture(str(VAL_PLOT), left, top, width=width)

    # --- SLIDE 18: SPH vs ODE & Re Sweep ---
    slide18 = prs.slides[17]
    print("[Slide Updater] Updating Slide 18 (SPH vs ODE Comparison)...")
    if COMP_PLOT.exists():
        left = Inches(6.5)
        top = Inches(1.4)
        width = Inches(6.5)
        slide18.shapes.add_picture(str(COMP_PLOT), left, top, width=width)

    # --- SLIDE 19: Dual-Mode Visual Renderer & Wake Vorticity ---
    slide19 = prs.slides[18]
    print("[Slide Updater] Updating Slide 19 (Sebastian Lague-Style Renderer)...")
    # Find any sample rendered frame
    vort_frames = list((FRAMES_DIR / "tumble_vorticity").glob("*.png"))
    speed_frames = list((FRAMES_DIR / "flutter_speed").glob("*.png"))
    
    sample_frame = None
    if vort_frames:
        sample_frame = vort_frames[min(5, len(vort_frames)-1)]
    elif speed_frames:
        sample_frame = speed_frames[min(5, len(speed_frames)-1)]

    if sample_frame and sample_frame.exists():
        left = Inches(7.0)
        top = Inches(1.4)
        width = Inches(5.8)
        slide19.shapes.add_picture(str(sample_frame), left, top, width=width)

    out_pptx = ROOT_DIR / "Autumn_Coin_IYPT_2026_Full_22Slides.pptx"
    prs.save(str(out_pptx))
    print(f"[Slide Updater] Successfully updated and saved {out_pptx}")

if __name__ == "__main__":
    update_slides()
